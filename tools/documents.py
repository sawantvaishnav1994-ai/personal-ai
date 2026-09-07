from pathlib import Path
from tools.registry import Tool,Risk

def register(reg,settings):
    root=(settings.data_dir/'workspace').resolve();root.mkdir(parents=True,exist_ok=True)
    def safe_path(raw,suffix):
        value=str(raw or '').strip() or f'untitled{suffix}'
        candidate=(root/value).resolve() if not Path(value).is_absolute() else Path(value).expanduser().resolve()
        if root not in candidate.parents and candidate!=root:raise ValueError(f'document output must stay inside Personal AI workspace: {root}')
        if candidate.suffix.lower()!=suffix:candidate=candidate.with_suffix(suffix)
        candidate.parent.mkdir(parents=True,exist_ok=True);return candidate
    def docx_create(p):
        from docx import Document
        path=safe_path(p.get('path'),'.docx');doc=Document()
        if p.get('title'):doc.add_heading(str(p['title']),0)
        for section in p.get('sections',[]):
            if section.get('heading'):doc.add_heading(str(section['heading']),level=min(max(int(section.get('level',1)),1),4))
            for para in section.get('paragraphs',[]):doc.add_paragraph(str(para))
        for para in p.get('paragraphs',[]):doc.add_paragraph(str(para))
        doc.save(path);return {'ok':True,'path':str(path),'kind':'docx'}
    def xlsx_create(p):
        from openpyxl import Workbook
        path=safe_path(p.get('path'),'.xlsx');wb=Workbook();sheets=p.get('sheets') or [{'name':'Sheet1','rows':p.get('rows',[])}]
        for index,item in enumerate(sheets):
            ws=wb.active if index==0 else wb.create_sheet();ws.title=str(item.get('name') or f'Sheet{index+1}')[:31]
            for row in item.get('rows',[]):ws.append(list(row))
        wb.save(path);return {'ok':True,'path':str(path),'kind':'xlsx'}
    def pptx_create(p):
        from pptx import Presentation
        path=safe_path(p.get('path'),'.pptx');prs=Presentation()
        for item in p.get('slides',[]):
            slide=prs.slides.add_slide(prs.slide_layouts[1]);slide.shapes.title.text=str(item.get('title',''));slide.placeholders[1].text=str(item.get('body',''))
        prs.save(path);return {'ok':True,'path':str(path),'kind':'pptx'}
    def pdf_create(p):
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,PageBreak
        path=safe_path(p.get('path'),'.pdf');styles=getSampleStyleSheet();story=[]
        if p.get('title'):story.extend([Paragraph(str(p['title']),styles['Title']),Spacer(1,14)])
        for section in p.get('sections',[]):
            if section.get('heading'):story.extend([Paragraph(str(section['heading']),styles['Heading2']),Spacer(1,6)])
            for para in section.get('paragraphs',[]):story.extend([Paragraph(str(para),styles['BodyText']),Spacer(1,8)])
            if section.get('page_break'):story.append(PageBreak())
        for para in p.get('paragraphs',[]):story.extend([Paragraph(str(para),styles['BodyText']),Spacer(1,8)])
        SimpleDocTemplate(str(path),pagesize=A4,title=str(p.get('title','Personal AI document'))).build(story or [Paragraph('',styles['BodyText'])]);return {'ok':True,'path':str(path),'kind':'pdf'}
    reg.register(Tool('create_docx','Create a Word document inside the Personal AI workspace; params: path,title,paragraphs,sections',docx_create,Risk.REVERSIBLE))
    reg.register(Tool('create_xlsx','Create a spreadsheet inside the Personal AI workspace; params: path,rows or sheets',xlsx_create,Risk.REVERSIBLE))
    reg.register(Tool('create_pptx','Create a presentation inside the Personal AI workspace; params: path,slides',pptx_create,Risk.REVERSIBLE))
    reg.register(Tool('create_pdf','Create a PDF inside the Personal AI workspace; params: path,title,paragraphs,sections',pdf_create,Risk.REVERSIBLE))
